"""National Bank: Canadian PCL and loans by product, capital, earnings.

long-comment-ok: disclosure notes, same format as rbc.py / cibc.py
What NBC's SFI discloses (checked FY2018-FY2026 packages):
  * Gross loans by geography x Basel category (Canada: residential mortgages incl. HELOCs,
    qualifying revolving retail, other retail, non-retail) -> gross_loans, direct
  * PCL by business segment, Stage 3 / Stages 1-2 / POCI -> pcl_total, ALLOCATED (decisions.md):
      Credit card (P&C)                 -> Qualifying revolving retail
      Personal Banking (P&C)            -> Residential mortgages / Other retail, split by each
                                           category's share of trailing-4Q impaired-loan PCL
                                           (total bank, by Basel borrower category)
      Wealth Management                 -> Other retail
      Commercial Banking (P&C, incl. CWB POCI from FQ2 2025) -> Non-retail
      Capital/Financial Markets         -> Non-retail x Canada's share of non-retail loans
                                           (Canada / Canada + U.S. + Europe)
      USSF&I (Credigy, ABA Bank)        -> excluded
  * Capital (CET1 capital, RWA, CET1 ratio) from the separate regulatory capital supplement
    (NBC_FY*_REGCAP.xlsx); RWA = CET1 capital RWA while the CVA phase-in applied.
  * FY2021-FY2024 "-revised" SFIs are partial restatements; full originals are *_SFIfull.xlsx.
"""
import re
from dataclasses import dataclass, field
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

BANK = "NBC"
PRODUCT_RE = {  # emitted label -> regex on the Basel category label (FY2018 uses singular / no hyphen)
    "Residential mortgages": r"Residential mortgages?",
    "Qualifying revolving retail": r"Qualifying revolving retail",
    "Other retail": r"Other retail",
    "Non-retail": r"Non[- ]retail",
}
PRODUCTS = list(PRODUCT_RE)
Q_RE = re.compile(r"^Q([1-4])$")


@dataclass
class Row:
    idx: int
    path: tuple  # carried labels of the label columns, footnote markers removed
    own: str  # labels written on this row (blank for subtotal rows)
    raw: list = field(repr=False)
    colmap: dict = field(repr=False)  # column -> (fiscal_year, fiscal_quarter)

    def vals(self, offset: int = 0) -> dict:
        """Values in the column `offset` places right of each quarter's first column."""
        out = {}
        for j, q in self.colmap.items():
            v = self.raw[j + offset] if j + offset < len(self.raw) else None
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                out[q] = float(v)
        return out

    @property
    def values(self) -> dict:
        return self.vals(0)


def _clean(s) -> str:
    if not isinstance(s, str):
        return ""
    s = re.sub(r"_x000D_|\s+", " ", s)
    s = re.sub(r"(\(\d+\))+", "", s)  # footnote markers like (1)(2)
    return s.strip().rstrip(":").strip()


def _colmap(rows, i) -> dict | None:
    """Quarter row i (Q1..Q4 cells) under a year row; years carry right until a text cell."""
    r = rows[i]
    if sum(bool(isinstance(c, str) and Q_RE.match(c.strip())) for c in r) < 3:
        return None
    above, yr, cols = rows[i - 1], None, {}
    for j, c in enumerate(r):
        y = above[j] if j < len(above) else None
        if isinstance(y, (int, float)) and 1990 < y < 2100:
            yr = int(y)
        elif isinstance(y, str) and y.strip():
            yr = None
        if yr and isinstance(c, str) and Q_RE.match(c.strip()):
            cols.setdefault((yr, int(c.strip()[1])), j)  # first column of each quarter
    return {j: q for q, j in cols.items()} if cols else None


def read_sheet(ws, label_cols: int = 4, max_rows: int = 300) -> list[Row]:
    rows = [list(r) for r in ws.iter_rows(max_row=max_rows, values_only=True)]
    out, colmap, path = [], None, [""] * label_cols
    for i, r in enumerate(rows):
        cm = _colmap(rows, i) if i else None
        if cm:
            colmap, path = cm, [""] * label_cols
            continue
        if colmap is None:
            continue
        labs = [_clean(c) for c in (r + [None] * label_cols)[:label_cols]]
        for k, lab in enumerate(labs):
            if lab:
                path[k] = lab
                path[k + 1:] = [""] * (label_cols - k - 1)
        out.append(Row(i + 1, tuple(p for p in path if p), " ".join(l for l in labs if l), r, colmap))
    return out


def load(path: Path) -> list[tuple[str, list[Row]]]:
    """(title, rows) per sheet; title is the upper-cased text of the first three rows."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        head = [c for r in ws.iter_rows(max_row=3, values_only=True) for c in r if isinstance(c, str)]
        out.append((" ".join(_clean(c) for c in head).upper(), read_sheet(ws)))
    return out


def _sheets(book, title: str, exclude: str | None = None, anywhere: bool = False) -> list[list[Row]]:
    title = title.upper()
    return [rows for t, rows in book
            if (title in t if anywhere else t.startswith(title)) and not (exclude and exclude in t)]


def _first(rows: list[Row], pattern: str, offset: int = 0, path0: str | None = None) -> dict:
    rx = re.compile(pattern, re.I)
    for r in rows:
        if path0 and (not r.path or not re.search(path0, r.path[0], re.I)):
            continue
        if rx.search(r.own) and r.vals(offset):
            return r.vals(offset)
    return {}


# --- per-package parsing -----------------------------------------------------

SEGMENTS = {  # key -> (path[0] pattern, path[1] pattern or None)
    "personal": (r"^Personal and Commercial", r"^(Personal Banking|Retail)$"),
    "card": (r"^Personal and Commercial", r"^Credit card"),
    "commercial": (r"^Personal and Commercial", r"^Commercial( Banking)?$"),
    "wealth": (r"^Wealth Management", None),
    "markets": (r"^(Financial|Capital) Markets", None),
}


def _kind(own: str) -> str:
    if not own:
        return "total"
    if re.search(r"POCI", own, re.I):
        return "poci"
    if re.search(r"Stage 3|Impaired", own, re.I):
        return "impaired"
    if re.search(r"Stages? 1 and 2|Performing", own, re.I):
        return "performing"
    return "other"


def parse_pcl(book) -> dict:
    """segment -> kind -> {quarter: value}."""
    out = {k: {} for k in SEGMENTS}
    for rows in _sheets(book, "PROVISIONS FOR CREDIT LOSSES"):
        for key, (p0, p1) in SEGMENTS.items():
            seg = out[key]
            for r in rows:
                if not r.values or not r.path or not re.search(p0, r.path[0], re.I):
                    continue
                if p1 and (len(r.path) < 2 or not re.search(p1, r.path[1], re.I)):
                    continue
                kind = _kind(r.own if not p1 else r.own.replace(r.path[1], "").strip())
                if kind in seg and kind != "poci":
                    continue  # first row of each kind wins (later blank rows are grand totals)
                if kind == "poci" and "poci" in seg:
                    continue
                seg[kind] = r.values
    return out


def parse_loans(book) -> tuple[dict, dict]:
    """Canadian gross loans by Basel category; non-retail by region (for the markets share)."""
    loans, nonretail = {}, {}
    for rows in _sheets(book, "GEOGRAPHIC DISTRIBUTION"):
        for r in rows:
            if not r.values or len(r.path) < 2 or not r.own:
                continue
            region, cat = r.path[0], r.own
            for prod, rx in PRODUCT_RE.items():
                if re.fullmatch(rx, cat, re.I) and region == "Canada":
                    loans.setdefault(prod, {}).update(r.values)
            if re.fullmatch(PRODUCT_RE["Non-retail"], cat, re.I) and region in ("Canada", "United States", "Europe"):
                nonretail.setdefault(region, {}).update(r.values)
    return loans, nonretail


def parse_impaired_pcl_by_category(book) -> dict:
    """Total-bank PCL on impaired loans by Basel category (4th column of each quarter)."""
    out = {}
    for rows in _sheets(book, "BORROWER CATEGORY", anywhere=True):
        for prod in ("Residential mortgages", "Other retail"):
            v = _first(rows, rf"^{PRODUCT_RE[prod]}$", offset=3)
            if v:
                out.setdefault(prod, {}).update(v)
    return out


def parse_income(book) -> dict:
    cons = _sheets(book, "CONSOLIDATED RESULTS")
    det = _sheets(book, "DETAILED INFORMATION ON INCOME", exclude="TAXABLE EQUIVALENT")
    out = {}
    for rows in cons or det:
        pretax = _first(rows, r"^Income before income taxes$")
        if pretax:
            out = {"pretax": pretax, "pcl": _first(rows, r"^Provisions for credit losses$")}
            break
    for rows in _sheets(book, "DETAILED INFORMATION ON INCOME") + _sheets(book, "CONSOLIDATED RESULTS"):
        d = _first(rows, r"^Dividends on common shares$")
        if d:
            out["divs"] = d
            break
    return out


def parse_capital(book) -> dict:
    cet1, rwa_cet1, rwa_total, ratio = {}, {}, {}, {}
    for _, rows in book:
        for r in rows:
            if not r.values:
                continue
            if re.search(r"^Common Equity Tier 1 capital \(CET1\)$", r.own) and not cet1:
                cet1 = r.values
            elif re.search(r"Common Equity Tier 1 Capital RWA", r.own, re.I) and not rwa_cet1:
                rwa_cet1 = r.values
            elif re.fullmatch(r"Total risk-weighted assets", r.own, re.I) and not rwa_total:
                rwa_total = r.values
            elif re.search(r"^(\d+\w?\s+)?Common Equity Tier 1 \(as a percentage of risk.weighted assets\)",
                           r.own, re.I) and not ratio:
                ratio = r.values
    rwa = dict(rwa_total)
    rwa.update(rwa_cet1)  # CET1 RWA where disclosed (CVA phase-in years)
    return {"cet1": cet1, "rwa": rwa, "ratio": ratio}


# --- assembly ----------------------------------------------------------------

def _merge(parsed: list[dict], getter) -> dict:
    """Combine one series across packages; the newest package wins for each quarter."""
    out = {}
    for p in parsed:  # newest first
        for q, v in getter(p).items():
            out.setdefault(q, v)
    return out


def _order(f: Path) -> tuple:
    """Newest first; within a year the full package outranks a partial '-revised' one,
    since the full package carries every table (the revised one only restated segments)."""
    m = re.search(r"FY(\d{4})Q(\d)", f.name)
    return (int(m.group(1)), int(m.group(2)), "full" in f.name)


def extract(package_dir: Path, start=(2018, 1)) -> tuple[pd.DataFrame, pd.DataFrame]:
    sfis = sorted(package_dir.glob(f"{BANK}_FY*_SFI*.xlsx"), key=_order, reverse=True)
    regcaps = sorted(package_dir.glob(f"{BANK}_FY*_REGCAP.xlsx"), key=_order, reverse=True)
    parsed = []
    for f in sfis:
        book = load(f)
        loans, nonretail = parse_loans(book)
        parsed.append({"file": f.name, "pcl": parse_pcl(book), "loans": loans, "nonretail": nonretail,
                       "imp_cat": parse_impaired_pcl_by_category(book), "income": parse_income(book)})
    caps = [parse_capital(load(f)) | {"file": f.name} for f in regcaps]

    pcl = {s: {k: _merge(parsed, lambda p, s=s, k=k: p["pcl"][s].get(k, {}))
               for k in ("total", "impaired", "performing", "poci")} for s in SEGMENTS}
    loans = {k: _merge(parsed, lambda p, k=k: p["loans"].get(k, {})) for k in PRODUCTS}
    nonretail = {k: _merge(parsed, lambda p, k=k: p["nonretail"].get(k, {}))
                 for k in ("Canada", "United States", "Europe")}
    imp_cat = {k: _merge(parsed, lambda p, k=k: p["imp_cat"].get(k, {}))
               for k in ("Residential mortgages", "Other retail")}
    income = {k: _merge(parsed, lambda p, k=k: p["income"].get(k, {})) for k in ("pretax", "pcl", "divs")}
    cap = {k: _merge(caps, lambda p, k=k: p[k]) for k in ("cet1", "rwa", "ratio")}
    src = {}
    for p in reversed(parsed):
        for q in p["loans"].get("Non-retail", {}):
            src[q] = p["file"]

    quarters = sorted(q for q in loans["Non-retail"] if q >= start and q <= max(pcl["personal"]["total"]))

    # Personal Banking split: trailing-4Q share of impaired-loan PCL, mortgages vs other retail
    cat = pd.DataFrame(imp_cat).sort_index().clip(lower=0).rolling(4, min_periods=1).sum()
    w_mort = (cat["Residential mortgages"] / cat.sum(axis=1)).fillna(0.0).to_dict()

    def mortgage_weight(q) -> float:
        prior = [k for k in w_mort if k <= q]
        return w_mort[max(prior)] if prior else w_mort[min(w_mort)]

    def seg(s, kind, q):
        return pcl[s][kind].get(q, np.nan)

    def components(s, q):
        imp = seg(s, "impaired", q) + (seg(s, "poci", q) if q in pcl[s]["poci"] else 0.0)
        perf = seg(s, "performing", q)
        return perf, imp

    rows = []
    for q in quarters:
        nr = nonretail
        ca_share = nr["Canada"][q] / (nr["Canada"][q] + nr["United States"].get(q, 0) + nr["Europe"].get(q, 0))
        w = mortgage_weight(q)
        pb_perf, pb_imp = components("personal", q)
        cc_perf, cc_imp = components("card", q)
        wm_perf, wm_imp = components("wealth", q)
        cb_perf, cb_imp = components("commercial", q)
        cm_perf, cm_imp = components("markets", q)
        parts = {
            "Residential mortgages": (pb_perf * w, pb_imp * w),
            "Qualifying revolving retail": (cc_perf, cc_imp),
            "Other retail": (pb_perf * (1 - w) + wm_perf, pb_imp * (1 - w) + wm_imp),
            "Non-retail": (cb_perf + cm_perf * ca_share, cb_imp + cm_imp * ca_share),
        }
        for prod, (perf, imp) in parts.items():
            rows.append({
                "bank": BANK, "fiscal_year": q[0], "fiscal_quarter": q[1], "segment_label": prod,
                "pcl_total": round(perf + imp, 2), "pcl_performing": round(perf, 2),
                "pcl_impaired": round(imp, 2), "gross_loans": loans[prod].get(q, np.nan),
                "source": src.get(q, ""),
                "notes": "Loans direct (Canada, Basel category); PCL allocated from segments (decisions.md)",
            })
    pcl_df = pd.DataFrame(rows)

    cap_df = pd.DataFrame([{
        "bank": BANK, "fiscal_year": q[0], "fiscal_quarter": q[1],
        "cet1_capital": cap["cet1"].get(q), "rwa": cap["rwa"].get(q), "cet1_ratio": cap["ratio"].get(q),
        "ppt_earnings": income["pretax"].get(q, np.nan) + income["pcl"].get(q, np.nan),
        "pcl_total_bank": income["pcl"].get(q), "common_dividends": income["divs"].get(q),
        "source": src.get(q, ""),
        "notes": "Capital from regulatory capital supplement; RWA = CET1 capital RWA during CVA phase-in",
    } for q in quarters])
    return pcl_df, cap_df
